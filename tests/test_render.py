import json
from pathlib import Path

from pptx import Presentation

from informe_mensual.periodos import Periodo
from informe_mensual.render import nombre_archivo, renderizar
from informe_mensual.spec import InformeMensual

from .fixtures import INFORME_VARIADO

RAIZ = Path(__file__).resolve().parent.parent


def _textos(pptx_path):
    prs = Presentation(str(pptx_path))
    return prs, [" ".join(sh.text_frame.text for sh in s.shapes if sh.has_text_frame) for s in prs.slides]


def test_ejemplo_septiembre_reproduce_estructura_del_mockup(tmp_path):
    informe = InformeMensual.model_validate(json.loads((RAIZ / "examples" / "2026-09.json").read_text()))
    destino = renderizar(informe, Periodo(2026, 9), tmp_path / "sep.pptx")
    prs, textos = _textos(destino)
    assert len(prs.slides) == 17  # mismo largo que el mockup de septiembre
    assert "Septiembre 2026" in textos[0]
    assert "215 horas totales trabajadas" in textos[0]
    assert "01 / 17" in textos[0] and "17 / 17" in textos[-1]
    assert "98 hs" in textos[1] and "46% del total" in textos[1]
    assert "▼ 37,00% vs ago/26" in textos[3]
    assert "¡Gracias!" in textos[-1]


def test_tipos_variados(tmp_path):
    destino = renderizar(INFORME_VARIADO, Periodo(2026, 7), tmp_path / "jul.pptx")
    prs, textos = _textos(destino)
    # portada + áreas + 1 sección + 5 diapositivas + cierre
    assert len(prs.slides) == 9
    assert "Julio 2026" in textos[0]
    tabla = [sh for sh in prs.slides[4].shapes if sh.has_table][0].table
    assert tabla.cell(0, 0).text == "KPI (Instagram)"
    assert tabla.cell(5, 3).text == "+570"


def test_nombre_archivo():
    assert nombre_archivo(Periodo(2026, 9), "Confort Iván") == "[Confort Iván] Septiembre 2026_ModusRetis.pptx"
