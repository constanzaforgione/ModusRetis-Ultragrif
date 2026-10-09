"""Dibuja la presentación mensual (.pptx) a partir de un `InformeMensual`.

Las medidas, colores y tipografías se tomaron de los mockups de julio, agosto
y septiembre 2026 (diapositiva 16:9 de 720 x 405 pt). Se parte de
`assets/plantilla_base.pptx`, que conserva el tema y las fuentes embebidas del
mockup, y cada diapositiva se arma con formas nativas editables.
"""
from __future__ import annotations

import math
from pathlib import Path
from typing import Optional, Sequence

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Pt

from . import config
from .periodos import Periodo
from .spec import (
    ColumnaLista,
    DiapoAds,
    DiapoDosColumnas,
    DiapoGruposMetricas,
    DiapoKpisLista,
    DiapoKpisTexto,
    DiapoListaNumerada,
    DiapoPlataforma,
    DiapoReputacion,
    DiapoTabla,
    DiapoTextoVinietas,
    DiapoVisionGeneral,
    InformeMensual,
    Kpi,
    Variacion,
)

ASSETS = Path(__file__).resolve().parent / "assets"

# Paleta del mockup
AZUL = "283B88"
AZUL_OSCURO = "1A2D6E"
CELESTE = "25BCD0"
CELESTE_CLARO = "B8E8F5"
FONDO = "FAFAFC"
BLANCO = "FFFFFF"
TEXTO = "3A3A45"
GRIS = "676779"
BORDE = "E9E9F4"
RECUADRO = "EFF2FF"
FILA_ALTERNA = "F3F5FC"
TONOS = {"positivo": "1E9E63", "negativo": "D14343", "advertencia": "BF9000", "neutral": "999999"}

TITULOS = "Roboto"
CUERPO = "Calibri"
MARCA = "Plus Jakarta Sans"

ANCHO, ALTO = 720.0, 405.0
MARGEN = 39.6
ANCHO_UTIL = 640.8


# ---------------------------------------------------------------------------
# Primitivas
# ---------------------------------------------------------------------------
def _rgb(hex_: str) -> RGBColor:
    return RGBColor.from_string(hex_)


def _alto_texto(texto: str, ancho: float, tam: float, factor: float, interlineado: float) -> float:
    lineas = 0
    for parrafo in texto.split("\n"):
        caracteres_por_linea = max(1.0, ancho / (tam * factor))
        lineas += max(1, math.ceil(len(parrafo) / caracteres_por_linea))
    return lineas * tam * interlineado


def ajustar_tamanio(
    textos: Sequence[str], ancho: float, alto: float, tam: float, minimo: float, factor: float = 0.5,
    interlineado: float = 1.2, espacio_parrafos: float = 0.0,
) -> float:
    """Achica la letra hasta que el texto estimado entre en la caja (heurística por ancho medio de caracter)."""
    while tam > minimo:
        total = sum(_alto_texto(t, ancho, tam, factor, interlineado) for t in textos)
        total += espacio_parrafos * tam * max(0, len(textos) - 1)
        if total <= alto:
            break
        tam -= 0.25
    return max(tam, minimo)


def caja_texto(
    slide, x, y, w, h, partes, *, tam: float, color: str = TEXTO, fuente: str = CUERPO, negrita: bool = False,
    cursiva: bool = False, alineacion=PP_ALIGN.LEFT, anclaje=MSO_ANCHOR.MIDDLE, espacio_despues: float = 0,
    interlineado: Optional[float] = None, subrayado: bool = False,
):
    """Crea un cuadro de texto sin márgenes internos.

    `partes` puede ser un string (un párrafo por cada '\\n'), o una lista de
    párrafos donde cada párrafo es una lista de (texto, overrides) para mezclar
    colores dentro de una misma línea.
    """
    tb = slide.shapes.add_textbox(Pt(x), Pt(y), Pt(w), Pt(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anclaje

    if isinstance(partes, str):
        parrafos = [[(linea, {})] for linea in partes.split("\n")]
    else:
        parrafos = partes

    for i, runs in enumerate(parrafos):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = alineacion
        if espacio_despues and i < len(parrafos) - 1:
            p.space_after = Pt(espacio_despues)
        if interlineado:
            p.line_spacing = interlineado
        for texto, extra in runs:
            r = p.add_run()
            r.text = texto
            f = r.font
            f.name = extra.get("fuente", fuente)
            f.size = Pt(extra.get("tam", tam))
            f.bold = extra.get("negrita", negrita)
            f.italic = extra.get("cursiva", cursiva)
            f.underline = extra.get("subrayado", subrayado)
            f.color.rgb = _rgb(extra.get("color", color))
    return tb


def tarjeta(slide, x, y, w, h, *, relleno: str = BLANCO, borde: str = BORDE, radio: float = 0.0457, sombra: bool = True):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Pt(x), Pt(y), Pt(w), Pt(h))
    _sin_estilo_tema(shp)
    shp.adjustments[0] = radio
    shp.fill.solid()
    shp.fill.fore_color.rgb = _rgb(relleno)
    if borde:
        shp.line.color.rgb = _rgb(borde)
        shp.line.width = Pt(1)
    else:
        shp.line.fill.background()
    if sombra:
        sp_pr = shp._element.spPr
        efectos = etree.SubElement(sp_pr, qn("a:effectLst"))
        sombra_el = etree.SubElement(
            efectos, qn("a:outerShdw"),
            blurRad="101600", dist="25400", dir="5400000", algn="bl", rotWithShape="0",
        )
        color = etree.SubElement(sombra_el, qn("a:srgbClr"), val="213478")
        etree.SubElement(color, qn("a:alpha"), val="12157")
    shp.text_frame.text = ""
    return shp


def punto(slide, x, y, tam: float = 7.2, color: str = CELESTE):
    shp = slide.shapes.add_shape(MSO_SHAPE.OVAL, Pt(x), Pt(y), Pt(tam), Pt(tam))
    _sin_estilo_tema(shp)
    shp.fill.solid()
    shp.fill.fore_color.rgb = _rgb(color)
    shp.line.fill.background()
    return shp


def linea(slide, x1, y1, x2, y2, color: str = BORDE, grosor: float = 1.0):
    ln = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Pt(x1), Pt(y1), Pt(x2), Pt(y2))
    _sin_estilo_tema(ln)
    ln.line.color.rgb = _rgb(color)
    ln.line.width = Pt(grosor)
    return ln


def imagen(slide, nombre: str, x, y, w, h):
    return slide.shapes.add_picture(str(ASSETS / nombre), Pt(x), Pt(y), Pt(w), Pt(h))


def _sin_estilo_tema(shp):
    """Quita el estilo heredado del tema (sombra y línea gruesa por defecto)."""
    estilo = shp._element.find(qn("p:style"))
    if estilo is not None:
        shp._element.remove(estilo)


def _fondo(slide, color: str):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = _rgb(color)


def _texto_variacion(v: Variacion) -> str:
    flecha = {"sube": "▲ ", "baja": "▼ ", "igual": ""}[v.direccion]
    return flecha + v.texto


# ---------------------------------------------------------------------------
# Componentes reutilizables
# ---------------------------------------------------------------------------
class Renderizador:
    def __init__(self, informe: InformeMensual, periodo: Periodo, cliente: str = config.CLIENTE):
        self.informe = informe
        self.periodo = periodo
        self.cliente = cliente
        self.prs = Presentation(str(ASSETS / "plantilla_base.pptx"))
        self.layout = self._layout_vacio()
        self.total = 0

    def _layout_vacio(self):
        for master in self.prs.slide_masters:
            for layout in master.slide_layouts:
                if len(layout.placeholders) == 0:
                    return layout
        return self.prs.slide_layouts[-1]

    # -- piezas comunes -----------------------------------------------------
    def _pie(self, slide, numero: int, oscuro: bool):
        color = CELESTE_CLARO if oscuro else GRIS
        caja_texto(slide, 36, 374.4, 108, 18, f"{numero:02d} / {self.total:02d}", tam=9, color=color, negrita=True)
        caja_texto(
            slide, 331.2, 374.4, 352.8, 18,
            f"{config.AGENCIA} x {self.cliente}  ·  Resumen {self.periodo.nombre}",
            tam=9, color=color, alineacion=PP_ALIGN.RIGHT,
        )

    def _nueva(self, oscuro: bool = False):
        slide = self.prs.slides.add_slide(self.layout)
        _fondo(slide, AZUL if oscuro else FONDO)
        if oscuro:
            imagen(slide, "destello.png", 180, -21.6, 360, 144)
        return slide

    def _titulo(self, slide, titulo: str, destacado: str = "", subtitulo: str = ""):
        partes = [(titulo, {})]
        if destacado:
            if titulo and not titulo.endswith(" ") and not destacado.startswith(" "):
                partes[0] = (titulo + " ", {})
            partes.append((destacado, {"color": CELESTE}))
        largo = len(titulo) + len(destacado)
        tam = 26 if largo <= 36 else 22
        caja_texto(slide, MARGEN, 36, 648, 39.6, [partes], tam=tam, fuente=TITULOS, negrita=True, color=AZUL)
        if subtitulo:
            caja_texto(slide, MARGEN, 72, 648, 21.6, subtitulo, tam=12, color=GRIS)

    def _kpi(self, slide, x, y, w, kpi: Kpi, tam_valor: float, *, sep_etiqueta: float = 30.3, sep_var: float = 47.6,
             alineacion=PP_ALIGN.LEFT, color_valor: str = AZUL):
        tam = min(tam_valor, ajustar_tamanio([kpi.valor], w, tam_valor * 1.3, tam_valor, 12, factor=0.62))
        caja_texto(slide, x, y, w, 36, kpi.valor, tam=tam, fuente=TITULOS, negrita=True, color=color_valor, alineacion=alineacion)
        caja_texto(slide, x, y + sep_etiqueta, w, 18, kpi.etiqueta.upper(), tam=8.5, color=GRIS, alineacion=alineacion)
        if kpi.variacion:
            caja_texto(
                slide, x, y + sep_var, w, 15.8, _texto_variacion(kpi.variacion), tam=9.5, negrita=True,
                color=TONOS[kpi.variacion.tono], alineacion=alineacion,
            )

    def _parrafos(self, slide, x, y, w, h, parrafos: Sequence[str], tam: float, minimo: float = 8.5,
                  color: str = TEXTO, cursiva: bool = False):
        parrafos = [p for p in parrafos if p.strip()]
        if not parrafos:
            return
        tam = ajustar_tamanio(parrafos, w, h, tam, minimo, interlineado=1.35, espacio_parrafos=1.2)
        caja_texto(
            slide, x, y, w, h, "\n".join(parrafos), tam=tam, color=color, anclaje=MSO_ANCHOR.TOP,
            espacio_despues=tam * 1.2, interlineado=1.15, cursiva=cursiva,
        )

    def _lista(self, slide, x, y, w, alto_total, items: Sequence[str], tam: float, *, pitch_max: float = 61.2,
               color: str = TEXTO):
        if not items:
            return
        pitch = min(pitch_max, alto_total / len(items))
        tam = ajustar_tamanio(items, w - 15.8, pitch * len(items), tam, 7.5)
        for i, item in enumerate(items):
            yi = y + i * pitch
            punto(slide, x, yi + pitch / 2 - 3.6)
            caja_texto(slide, x + 15.8, yi, w - 15.8, pitch, item, tam=tam, color=color)

    # -- diapositivas fijas --------------------------------------------------
    def portada(self, numero: int):
        s = self._nueva(oscuro=True)
        tarjeta(s, 28, 30.9, 64.4, 63.4, borde=BLANCO, radio=0.5, sombra=False)
        imagen(s, "logo.png", 39.6, 36.9, 41.2, 51.5)
        caja_texto(s, MARGEN, 126, 504, 21.6, "Ecommerce  ·  Análisis mensual", tam=13, color=CELESTE_CLARO, negrita=True)
        caja_texto(s, MARGEN, 149.8, 655.2, 63.4, "Resumen mensual", tam=52, fuente=TITULOS, negrita=True, color=BLANCO)
        caja_texto(s, MARGEN, 210.2, 655.2, 63.4, self.periodo.nombre, tam=52, fuente=TITULOS, negrita=True, color=CELESTE)
        caja_texto(s, MARGEN, 284.4, 540, 23, self.informe.canales, tam=12.5, color=CELESTE_CLARO)
        horas = _formato_horas(self.informe.horas_totales)
        caja_texto(
            s, MARGEN, 311, 600, 20.1,
            f"Período: {self.periodo.rango}   ·   {horas} horas totales trabajadas", tam=11, color=CELESTE_CLARO,
        )
        self._pie(s, numero, oscuro=True)

    def areas(self, numero: int):
        s = self._nueva()
        total = self.informe.horas_totales or sum(a.horas for a in self.informe.areas)
        subtitulo = self.informe.subtitulo_areas or (
            f"Distribución de horas por área de trabajo: {_formato_horas(total)} hs totales"
        )
        self._titulo(s, "Lo que hicimos ", f"en {self.periodo.nombre_mes}", subtitulo)
        areas = self.informe.areas
        n = len(areas)
        cols = 2 if n <= 4 else 3
        filas = math.ceil(n / cols)
        sep = 14.4
        w = (ANCHO_UTIL - sep * (cols - 1)) / cols
        h = min(126.0, (270 - 18 * (filas - 1)) / filas)
        for i, area in enumerate(areas):
            x = MARGEN + (i % cols) * (w + sep)
            y = 96.8 + (i // cols) * (h + 18)
            tarjeta(s, x, y, w, h)
            px, iw = x + 18, w - 36
            pct = round(area.horas / total * 100) if total else 0
            caja_texto(s, px, y + 13, iw, 36, f"{_formato_horas(area.horas)} hs", tam=28 if cols == 2 else 24,
                       fuente=TITULOS, negrita=True, color=AZUL)
            caja_texto(s, px, y + 49, iw, 18, f"{pct}% del total", tam=9.5, negrita=True, color=CELESTE)
            caja_texto(s, px, y + 69, iw, 21.6, area.nombre, tam=13.5 if cols == 2 else 12, negrita=True, color=TEXTO)
            tam_desc = ajustar_tamanio([area.descripcion], iw, h - 93, 9.2, 7)
            caja_texto(s, px, y + 91.5, iw, h - 93, area.descripcion, tam=tam_desc, color=GRIS, anclaje=MSO_ANCHOR.TOP)
        self._pie(s, numero, oscuro=False)

    def seccion(self, numero: int, indice: int, titulo: str, subtitulo: str):
        s = self._nueva(oscuro=True)
        caja_texto(s, MARGEN, 28.8, 252, 25.2, "MODUS RETIS", tam=12, fuente=MARCA, negrita=True, color=BLANCO)
        caja_texto(s, MARGEN, 144, 180, 187.2, str(indice), tam=220, fuente=TITULOS, negrita=True, color=AZUL_OSCURO)
        tam = 46 if len(titulo) <= 32 else 38
        caja_texto(s, 259.2, 169.2, 432, 64.8, titulo, tam=tam, fuente=TITULOS, negrita=True, color=BLANCO)
        caja_texto(s, 262.8, 251.4, 417.6, 50.4, subtitulo, tam=13, color=CELESTE_CLARO)
        self._pie(s, numero, oscuro=True)

    def cierre(self, numero: int):
        s = self._nueva(oscuro=True)
        caja_texto(s, MARGEN, 28.8, 252, 25.2, "MODUS RETIS", tam=12, fuente=MARCA, negrita=True, color=BLANCO)
        tarjeta(s, 292.9, 109.7, 134.1, 132, borde=BLANCO, radio=0.5, sombra=False)
        imagen(s, "logo.png", 317.1, 122.1, 85.8, 107.2)
        caja_texto(s, 0, 243.1, ANCHO, 49, config.AGENCIA, tam=38, fuente=TITULOS, negrita=True, color=BLANCO,
                   alineacion=PP_ALIGN.CENTER)
        caja_texto(s, 0, 293.5, ANCHO, 21.6, f"Ecommerce · {self.periodo.nombre}", tam=13, color=CELESTE_CLARO,
                   alineacion=PP_ALIGN.CENTER)
        linea(s, 309.6, 333.1, 410.4, 333.1, color=CELESTE, grosor=1.5)
        caja_texto(s, 0, 343.9, ANCHO, 36, "¡Gracias!", tam=18, fuente=TITULOS, color=BLANCO, alineacion=PP_ALIGN.CENTER)
        self._pie(s, numero, oscuro=True)

    # -- diapositivas de contenido ------------------------------------------
    def kpis_texto(self, s, d: DiapoKpisTexto):
        n = max(1, len(d.kpis))
        w = ANCHO_UTIL / n
        for i, k in enumerate(d.kpis):
            self._kpi(s, MARGEN + i * w, 111.6, w - 6, k, 26 if n <= 4 else 24)
        linea(s, MARGEN, 198, MARGEN + ANCHO_UTIL, 198)
        self._parrafos(s, MARGEN, 212.4, ANCHO_UTIL, 150, d.parrafos, 11.3)

    def reputacion(self, s, d: DiapoReputacion):
        y_metricas = 110.0
        if d.destacado:
            tarjeta(s, MARGEN, 83.5, ANCHO_UTIL, 72)
            x_texto = 57.6
            if d.destacado.icono == "insignia":
                imagen(s, "insignia.png", 50.2, 101.5, 36, 36)
                x_texto = 97.2
            caja_texto(s, x_texto, 92.9, 560, 25.2, d.destacado.titulo, tam=15, fuente=TITULOS, negrita=True, color=AZUL)
            caja_texto(s, x_texto, 116.6, 666 - x_texto, 32.4, d.destacado.texto, tam=10.5, color=GRIS,
                       anclaje=MSO_ANCHOR.TOP)
            y_metricas = 169.2
        n = max(1, len(d.metricas))
        sep = 9.4
        w = min(110.0, (ANCHO_UTIL - sep * (n - 1)) / n)
        x0 = MARGEN + (ANCHO_UTIL - (w * n + sep * (n - 1))) / 2
        for i, m in enumerate(d.metricas):
            x = x0 + i * (w + sep)
            tarjeta(s, x, y_metricas, w, 108, radio=0.0656)
            tam = ajustar_tamanio([m.valor], w - 6, 32, 24, 14, factor=0.62)
            caja_texto(s, x, y_metricas + 21.6, w, 39.6, m.valor, tam=tam, fuente=TITULOS, negrita=True,
                       color=TONOS["negativo"] if m.alerta else CELESTE, alineacion=PP_ALIGN.CENTER)
            caja_texto(s, x + 3.6, y_metricas + 66.2, w - 7.2, 36, m.etiqueta, tam=7.6, color=GRIS,
                       alineacion=PP_ALIGN.CENTER)
        y_texto = y_metricas + 108 + 21.6
        self._parrafos(s, MARGEN, y_texto, ANCHO_UTIL, 368 - y_texto, d.parrafos, 10.3, minimo=8)

    def ads(self, s, d: DiapoAds):
        cols = 3 if d.roas else 4
        w_col = 129.6
        if d.roas:
            tarjeta(s, 493.2, 104.4, 187.2, 129.6)
            caja_texto(s, 507.6, 113.8, 158.4, 18, d.roas.etiqueta.upper() or "ROAS", tam=10, negrita=True, color=GRIS)
            caja_texto(s, 504, 133.2, 165.6, 43.2, d.roas.valor, tam=34, fuente=TITULOS, negrita=True, color=AZUL)
            if d.roas.variacion:
                caja_texto(s, 507.6, 180, 165.6, 18, _texto_variacion(d.roas.variacion), tam=10.5, negrita=True,
                           color=TONOS[d.roas.variacion.tono])
            if d.roas_explicacion:
                caja_texto(s, 507.6, 200.2, 165.6, 28.8, d.roas_explicacion, tam=8, color=GRIS)
        else:
            w_col = ANCHO_UTIL / cols
        filas = math.ceil(len(d.kpis) / cols)
        pitch = 78.0 if filas <= 3 else 60.0
        for i, k in enumerate(d.kpis):
            x = MARGEN + (i % cols) * w_col
            y = 104.4 + (i // cols) * pitch
            self._kpi(s, x, y, w_col - 4, k, 18, sep_etiqueta=23.5, sep_var=39.6)
        y_nota = max(330.0, 104.4 + filas * pitch - 6)
        self._parrafos(s, MARGEN, y_nota, ANCHO_UTIL, 368 - y_nota, [d.nota], 9.2, minimo=7.5)

    def grupos_metricas(self, s, d: DiapoGruposMetricas):
        y = 104.0
        total_filas = sum(math.ceil(len(g.kpis) / 4) for g in d.grupos)
        pitch = 59.4 if total_filas + len(d.grupos) <= 4 else 50.0
        for g in d.grupos:
            caja_texto(s, MARGEN, y, 360, 25.2, g.nombre.upper(), tam=13 if len(d.grupos) > 1 else 15,
                       fuente=TITULOS, negrita=True, color=AZUL, subrayado=True)
            y += 26
            for i, k in enumerate(g.kpis):
                x = MARGEN + (i % 4) * 160.2
                yi = y + (i // 4) * pitch
                self._kpi(s, x, yi, 150, k, 21, sep_etiqueta=25.3, sep_var=41)
            y += math.ceil(len(g.kpis) / 4) * pitch + 4
        if d.nota:
            y_linea = max(y + 4, 230.0)
            linea(s, MARGEN, y_linea, MARGEN + ANCHO_UTIL, y_linea)
            self._parrafos(s, MARGEN, y_linea + 14, ANCHO_UTIL, 366 - (y_linea + 14), [d.nota], 11, minimo=8)

    def vision_general(self, s, d: DiapoVisionGeneral):
        n = max(1, len(d.kpis))
        w = (ANCHO_UTIL - 25.2) / n
        for i, k in enumerate(d.kpis):
            self._kpi(s, 64.8 + i * w, 136.8, w - 14.4, k, 30, sep_etiqueta=34.3, sep_var=51.6)
        linea(s, MARGEN, 241.2, MARGEN + ANCHO_UTIL, 241.2)
        tarjeta(s, MARGEN, 259.2, ANCHO_UTIL, 97.2, relleno=RECUADRO, sombra=False)
        caja_texto(s, 57.6, 272.2, 597.6, 21.6, d.recuadro.titulo, tam=13, fuente=TITULOS, negrita=True, color=AZUL)
        tam = ajustar_tamanio([d.recuadro.texto], 597.6, 54, 10.5, 8)
        caja_texto(s, 57.6, 295.2, 597.6, 54, d.recuadro.texto, tam=tam, anclaje=MSO_ANCHOR.TOP)

    def kpis_lista(self, s, d: DiapoKpisLista):
        anchos = [(MARGEN, 158.4), (216.0, 187.2)]
        for (x, w), k in zip(anchos, d.kpis[:2]):
            self._kpi(s, x, 108, w, k, 30, sep_etiqueta=34.3, sep_var=51.6)
        n = len(d.lista)
        alto_lista = 28.8 * n
        linea(s, 424.8, 108, 424.8, 108 + max(154.8, alto_lista + 30))
        caja_texto(s, 442.8, 108, 237.6, 21.6, d.lista_titulo, tam=12.5, fuente=TITULOS, negrita=True, color=AZUL)
        self._lista(s, 442.8, 132.5, 237.6, alto_lista, d.lista, 9.8, pitch_max=28.8)
        y_linea = max(252.0, 132.5 + alto_lista + 8)
        linea(s, MARGEN, y_linea, MARGEN + ANCHO_UTIL, y_linea)
        self._parrafos(s, MARGEN, y_linea + 14.4, ANCHO_UTIL, 366 - (y_linea + 14.4), [d.nota], 10.5, minimo=8)

    def plataforma(self, s, d: DiapoPlataforma):
        t = d.tarjeta
        tarjeta(s, MARGEN, 108, 313.2, 223.2)
        caja_texto(s, 57.6, 121, 277.2, 21.6, t.nombre, tam=14, fuente=TITULOS, negrita=True, color=AZUL)
        for i, k in enumerate(t.kpis[:2]):
            self._kpi(s, 57.6 + i * 136.8, 151.2, 129.6, k, 22, sep_etiqueta=26.3, sep_var=43.6)
        caja_texto(s, 57.6, 226.8, 277.2, 21.6, t.detalle, tam=9.8, color=GRIS)
        tam = ajustar_tamanio([t.nota], 277.2, 60, 9.5, 7.5)
        caja_texto(s, 57.6, 259.2, 277.2, 60, t.nota, tam=tam, color=GRIS, anclaje=MSO_ANCHOR.TOP)

    def _columna(self, s, x, w, col: ColumnaLista):
        caja_texto(s, x, 108, w, 21.6, col.titulo, tam=13, fuente=TITULOS, negrita=True, color=AZUL)
        if col.estilo == "recuadro":
            tarjeta(s, x, 136.8, w, 216, relleno=RECUADRO, sombra=False)
            items = col.items
            alto = 216 - 28.8
            pitch = alto / max(1, len(items))
            tam = ajustar_tamanio(items, w - 36, alto - 6 * len(items), 10, 7.5)
            for i, item in enumerate(items):
                ultimo = i == len(items) - 1 and len(items) > 1 and item.lstrip().startswith("→")
                caja_texto(s, x + 18, 151.2 + i * pitch, w - 36, pitch, item, tam=tam, negrita=ultimo,
                           color=AZUL if ultimo else TEXTO)
        else:
            tam = 9.8 if len(col.items) <= 4 else 8.6
            self._lista(s, x, 130, w, 232, col.items, tam)

    def dos_columnas(self, s, d: DiapoDosColumnas):
        self._columna(s, MARGEN, 302.4, d.izquierda)
        linea(s, 370.8, 108, 370.8, 360)
        self._columna(s, 392.4, 288, d.derecha)

    def lista_numerada(self, s, d: DiapoListaNumerada):
        n = max(1, len(d.items))
        pitch = min(52.6, 262 / n)
        h = pitch - 5.1
        for i, item in enumerate(d.items):
            y = 102.2 + i * pitch
            tarjeta(s, MARGEN, y, ANCHO_UTIL, h, radio=0.08)
            caja_texto(s, 54, y, 36, h, str(i + 1), tam=18, fuente=TITULOS, negrita=True, color=CELESTE)
            caja_texto(s, 100.8, y + 4.4, 565.2, 18.7, item.titulo, tam=11, fuente=TITULOS, negrita=True, color=AZUL)
            tam = ajustar_tamanio([item.texto], 565.2, h - 24, 8.8, 7)
            caja_texto(s, 100.8, y + 22.4, 565.2, h - 23.6, item.texto, tam=tam, anclaje=MSO_ANCHOR.TOP)

    def tabla(self, s, d: DiapoTabla):
        filas = len(d.filas) + 1
        cols = max(1, len(d.encabezados))
        alto_fila = 22.0
        tbl = s.shapes.add_table(filas, cols, Pt(MARGEN), Pt(104), Pt(ANCHO_UTIL), Pt(alto_fila * filas)).table
        for c, texto in enumerate(d.encabezados):
            self._celda(tbl.cell(0, c), texto, AZUL, BLANCO, negrita=True, centrar=c > 0)
        for r, fila in enumerate(d.filas, start=1):
            fondo = BLANCO if r % 2 else FILA_ALTERNA
            for c in range(cols):
                texto = fila[c] if c < len(fila) else ""
                self._celda(tbl.cell(r, c), texto, fondo, TEXTO, negrita=False, centrar=c > 0)
        if d.nota:
            y = 104 + alto_fila * filas + 18
            self._parrafos(s, MARGEN, y, ANCHO_UTIL, 366 - y, [d.nota], 9.5, minimo=7.5, cursiva=True)

    @staticmethod
    def _celda(celda, texto, fondo, color, *, negrita, centrar):
        celda.fill.solid()
        celda.fill.fore_color.rgb = _rgb(fondo)
        celda.margin_left = celda.margin_right = Pt(7.2)
        celda.margin_top = celda.margin_bottom = Pt(3.6)
        celda.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf = celda.text_frame
        tf.text = texto
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER if centrar else PP_ALIGN.LEFT
        for r in p.runs:
            r.font.name = CUERPO
            r.font.size = Pt(9)
            r.font.bold = negrita
            r.font.color.rgb = _rgb(color)

    def texto_vinietas(self, s, d: DiapoTextoVinietas):
        self._parrafos(s, MARGEN, 100, ANCHO_UTIL, 46, [d.intro], 10.3, minimo=8)
        mitad = math.ceil(len(d.items) / 2)
        columnas = [d.items[:mitad], d.items[mitad:]]
        alto = 26.0 * mitad
        for c, items in enumerate(columnas):
            self._lista(s, MARGEN + c * 327.6, 156, 300, alto, items, 8.6, pitch_max=26)
        if d.nota:
            y = 156 + alto + 14
            self._parrafos(s, MARGEN, y, ANCHO_UTIL, 366 - y, [d.nota], 8.6, minimo=7.5, color=GRIS, cursiva=True)

    # -- armado ---------------------------------------------------------------
    def construir(self) -> Presentation:
        plan: list = [("portada",), ("areas",)]
        for i, sec in enumerate(self.informe.secciones, start=1):
            plan.append(("seccion", i, sec))
            for d in sec.diapositivas:
                plan.append(("contenido", d))
        plan.append(("cierre",))
        self.total = len(plan)

        for numero, paso in enumerate(plan, start=1):
            tipo = paso[0]
            if tipo == "portada":
                self.portada(numero)
            elif tipo == "areas":
                self.areas(numero)
            elif tipo == "seccion":
                self.seccion(numero, paso[1], paso[2].titulo, paso[2].subtitulo)
            elif tipo == "cierre":
                self.cierre(numero)
            else:
                d = paso[1]
                s = self._nueva()
                self._titulo(s, d.titulo, d.titulo_destacado, d.subtitulo)
                getattr(self, d.tipo)(s, d)
                self._pie(s, numero, oscuro=False)
        return self.prs


def _formato_horas(horas: float) -> str:
    if float(horas).is_integer():
        return f"{int(horas)}"
    return f"{horas:.1f}".replace(".", ",")


def nombre_archivo(periodo: Periodo, cliente: str = config.CLIENTE) -> str:
    """Mismo formato que los mockups: '[Confort Iván] Septiembre 2026_ModusRetis.pptx'."""
    return f"[{cliente}] {periodo.nombre}_ModusRetis.pptx"


def renderizar(informe: InformeMensual, periodo: Periodo, destino: Path, cliente: str = config.CLIENTE) -> Path:
    prs = Renderizador(informe, periodo, cliente).construir()
    destino = Path(destino)
    destino.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(destino))
    return destino

